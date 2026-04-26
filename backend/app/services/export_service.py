import io
from typing import Any
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime


class ExportService:
    """Service for exporting analysis results to various formats."""

    def generate_xlsx(self, data: list[dict], job_title: str) -> tuple[bytes, str]:
        """
        Generate Excel file with analysis results.

        Returns:
            tuple: (file_content, filename)
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "分析结果"

        # Define headers
        headers = [
            "候选人姓名",
            "邮箱",
            "综合评分",
            "推荐等级",
            "技能匹配度",
            "经验匹配度",
            "教育背景",
            "项目相关性",
            "整体质量"
        ]

        # Header styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")

        # Write headers
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        # Write data
        for row_num, row_data in enumerate(data, 2):
            ws.cell(row=row_num, column=1).value = row_data.get("candidate_name", "")
            ws.cell(row=row_num, column=2).value = row_data.get("email", "")
            ws.cell(row=row_num, column=3).value = row_data.get("overall_score", 0)
            ws.cell(row=row_num, column=4).value = self._translate_recommendation(row_data.get("recommendation", ""))
            ws.cell(row=row_num, column=5).value = row_data.get("skill_match", 0)
            ws.cell(row=row_num, column=6).value = row_data.get("experience_match", 0)
            ws.cell(row=row_num, column=7).value = row_data.get("education", 0)
            ws.cell(row=row_num, column=8).value = row_data.get("project_relevance", 0)
            ws.cell(row=row_num, column=9).value = row_data.get("overall_quality", 0)

        # Auto-adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        file_content = output.getvalue()

        # Generate filename
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{job_title}_分析结果_{timestamp}.xlsx"

        return file_content, filename

    def generate_csv(self, data: list[dict], job_title: str) -> tuple[bytes, str]:
        """
        Generate CSV file with analysis results.

        Returns:
            tuple: (file_content, filename)
        """
        import csv

        output = io.StringIO()
        writer = csv.writer(output)

        # Write headers
        headers = [
            "候选人姓名",
            "邮箱",
            "综合评分",
            "推荐等级",
            "技能匹配度",
            "经验匹配度",
            "教育背景",
            "项目相关性",
            "整体质量"
        ]
        writer.writerow(headers)

        # Write data
        for row_data in data:
            writer.writerow([
                row_data.get("candidate_name", ""),
                row_data.get("email", ""),
                row_data.get("overall_score", 0),
                self._translate_recommendation(row_data.get("recommendation", "")),
                row_data.get("skill_match", 0),
                row_data.get("experience_match", 0),
                row_data.get("education", 0),
                row_data.get("project_relevance", 0),
                row_data.get("overall_quality", 0)
            ])

        # Get bytes content
        file_content = output.getvalue().encode("utf-8-sig")

        # Generate filename
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{job_title}_分析结果_{timestamp}.csv"

        return file_content, filename

    def _translate_recommendation(self, recommendation: str) -> str:
        """Translate recommendation enum to Chinese."""
        translation = {
            "strongly_recommended": "强烈推荐",
            "recommended": "推荐",
            "pending": "待定"
        }
        return translation.get(recommendation, recommendation)
