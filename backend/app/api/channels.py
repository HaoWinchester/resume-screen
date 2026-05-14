import csv
import html
import io
import re
import uuid
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import User, get_current_user
from app.database import get_db
from app.models import (
    AnalysisResult,
    AnalysisStatus,
    JobRequirement,
    JobRequirementStatus,
    ParseStatus,
    RecommendationLevel,
    Resume,
    ResumeFileType,
)
from app.schemas.channels import (
    ChannelCandidateImportItem,
    ChannelCandidateImportRequest,
    ChannelCandidateImportResponse,
)
from app.services.resume_parser import ResumeParser
from app.services.task_dispatcher import dispatch_resume_analysis

router = APIRouter(prefix="/channels", tags=["channels"])

PLATFORM_LABELS = {
    "boss_zhipin": "BOSS直聘",
    "liepin": "猎聘",
    "lagou": "拉勾",
    "email": "邮件投递",
    "other": "外部渠道",
}


def _normalize_imported_content(content: str, content_format: str) -> str:
    """Convert authorized exported text/HTML into plain resume text."""
    text = content.strip()
    if content_format == "html":
        text = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", text, flags=re.I | re.S)
        text = re.sub(r"</?(br|p|div|li|tr|section|article|h[1-6])[^>]*>", "\n", text, flags=re.I)
        text = re.sub(r"<[^>]+>", " ", text)

    text = html.unescape(text)
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n\s+", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _build_file_name(platform: str, candidate_name: str | None, resume_id: uuid.UUID) -> str:
    label = PLATFORM_LABELS.get(platform, "外部渠道")
    safe_name = candidate_name or "候选人"
    safe_name = re.sub(r"[\\/:*?\"<>|]+", "_", safe_name).strip() or "候选人"
    return f"{label}-{safe_name}-{str(resume_id)[:8]}.pdf"


def _pick_value(row: dict[str, str], *keys: str) -> str | None:
    for key in keys:
        value = row.get(key)
        if value and value.strip():
            return value.strip()
    return None


def _row_to_candidate(row: dict[str, str]) -> ChannelCandidateImportItem:
    normalized = {str(key).strip().lower(): str(value or "").strip() for key, value in row.items() if key}
    content = _pick_value(normalized, "content", "resume_text", "profile_text", "text", "简历内容", "候选人内容")
    if not content:
        content_parts = [
            f"{key}: {value}"
            for key, value in normalized.items()
            if value and key not in {"source_url", "url", "profile_url", "external_candidate_id", "id"}
        ]
        content = "\n".join(content_parts)

    content_format = _pick_value(normalized, "content_format", "format") or "text"
    if content_format not in {"text", "html"}:
        content_format = "text"

    return ChannelCandidateImportItem(
        candidate_name=_pick_value(normalized, "candidate_name", "name", "姓名", "候选人姓名"),
        source_url=_pick_value(normalized, "source_url", "url", "profile_url", "来源链接"),
        external_candidate_id=_pick_value(normalized, "external_candidate_id", "candidate_id", "id", "外部id"),
        content=content,
        content_format=content_format,
    )


async def _parse_candidate_file(file: UploadFile) -> list[ChannelCandidateImportItem]:
    filename = (file.filename or "").lower()
    raw = await file.read()
    if not raw:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={"error": {"code": "EMPTY_FILE", "message": "批量导入文件为空"}},
        )

    rows: list[dict[str, str]] = []
    if filename.endswith(".csv"):
        text = raw.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows = [dict(row) for row in reader]
    elif filename.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            rows = []
        else:
            headers = [str(value or "").strip() for value in values[0]]
            for record in values[1:]:
                rows.append({headers[index]: str(value or "") for index, value in enumerate(record) if index < len(headers)})
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={"error": {"code": "UNSUPPORTED_FILE", "message": "仅支持 CSV 或 XLSX 批量导入文件"}},
        )

    candidates = [_row_to_candidate(row) for row in rows if any(str(value or "").strip() for value in row.values())]
    if not candidates:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={"error": {"code": "NO_CANDIDATES", "message": "文件中没有可导入的候选人"}},
        )
    if len(candidates) > 500:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={"error": {"code": "TOO_MANY_CANDIDATES", "message": "单次最多导入 500 位候选人"}},
        )

    return candidates


async def _import_candidates(
    *,
    job_requirement_id: uuid.UUID,
    source_platform: str,
    candidates: list[ChannelCandidateImportItem],
    current_user: User,
    db: AsyncSession,
) -> dict:
    result = await db.execute(
        select(JobRequirement).where(
            JobRequirement.id == job_requirement_id,
            JobRequirement.company_id == current_user.company_id,
        )
    )
    job = result.scalar_one_or_none()
    if not job:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail={"error": {"code": "NOT_FOUND", "message": "岗位需求不存在"}},
        )

    if job.status != JobRequirementStatus.ACTIVE:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={"error": {"code": "INVALID_JOB_STATUS", "message": "仅活跃状态的岗位可导入候选人"}},
        )

    parser = ResumeParser()
    imported = []
    failed = []
    imported_at = datetime.now(timezone.utc)

    for item in candidates:
        try:
            raw_text = _normalize_imported_content(item.content, item.content_format)
            if len(raw_text) < 20:
                raise ValueError("候选人内容过短，无法解析")

            parsed_data = parser.extract_structured_data(raw_text)
            provided_name = parser.clean_candidate_name(item.candidate_name)
            if provided_name:
                parsed_data["name"] = provided_name

            parsed_data["source"] = {
                "type": "channel_import",
                "platform": source_platform,
                "platform_label": PLATFORM_LABELS.get(source_platform, "外部渠道"),
                "source_url": item.source_url,
                "external_candidate_id": item.external_candidate_id,
                "imported_at": imported_at.isoformat(),
            }
            parsed_data["raw_text"] = raw_text

            resume_id = uuid.uuid4()
            candidate_name = parser.clean_candidate_name(parsed_data.get("name"))
            resume = Resume(
                id=resume_id,
                job_requirement_id=job_requirement_id,
                file_name=_build_file_name(source_platform, candidate_name, resume_id),
                file_path=f"channel-import://{source_platform}/{resume_id}",
                # Keep the existing enum stable for deployed databases; imported
                # channel candidates are already parsed and do not need file parsing.
                file_type=ResumeFileType.PDF,
                file_size=len(raw_text.encode("utf-8")),
                parse_status=ParseStatus.SUCCESS,
                parse_error=None,
                parsed_data=parsed_data,
                candidate_name=candidate_name,
                candidate_email=parsed_data.get("email"),
                candidate_phone=parsed_data.get("phone"),
                uploaded_by=current_user.id,
            )
            db.add(resume)

            analysis = AnalysisResult(
                resume_id=resume_id,
                job_requirement_id=job_requirement_id,
                overall_score=0.0,
                recommendation=RecommendationLevel.PENDING,
                analysis_status=AnalysisStatus.PENDING,
            )
            db.add(analysis)

            await db.commit()

            dispatch_resume_analysis(str(resume_id), str(job_requirement_id))

            imported.append(
                {
                    "id": resume_id,
                    "resume_id": resume_id,
                    "candidate_name": candidate_name,
                    "source_platform": source_platform,
                    "parse_status": ParseStatus.SUCCESS.value,
                    "analysis_status": AnalysisStatus.PENDING.value,
                }
            )
        except Exception as exc:
            await db.rollback()
            failed.append(
                {
                    "candidate_name": item.candidate_name,
                    "source_url": item.source_url,
                    "error": str(exc),
                }
            )

    return {
        "job_requirement_id": job_requirement_id,
        "imported": imported,
        "failed": failed,
        "total_imported": len(imported),
        "total_failed": len(failed),
        "imported_at": imported_at,
    }


@router.post(
    "/candidates/import",
    response_model=ChannelCandidateImportResponse,
    status_code=status.HTTP_202_ACCEPTED,
)
async def import_channel_candidates(
    payload: ChannelCandidateImportRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """
    导入外部招聘渠道的授权候选人内容。

    这个接口不做未授权平台抓取，也不保存平台账号密码。它接收 HR 已获得授权
    的候选人文本/HTML 导出内容，复用当前简历解析与 AI 匹配分析链路。
    """
    return await _import_candidates(
        job_requirement_id=payload.job_requirement_id,
        source_platform=payload.source_platform,
        candidates=payload.candidates,
        current_user=current_user,
        db=db,
    )


@router.post(
    "/candidates/import-file",
    response_model=ChannelCandidateImportResponse,
    status_code=status.HTTP_202_ACCEPTED,
)
async def import_channel_candidates_file(
    job_requirement_id: uuid.UUID = Form(...),
    source_platform: str = Form("other"),
    consent_confirmed: bool = Form(...),
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """从 CSV/XLSX 批量导入已授权候选人资料。"""
    if not consent_confirmed:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={"error": {"code": "CONSENT_REQUIRED", "message": "必须确认候选人数据来源已获得授权"}},
        )
    if source_platform not in PLATFORM_LABELS:
        source_platform = "other"

    candidates = await _parse_candidate_file(file)
    return await _import_candidates(
        job_requirement_id=job_requirement_id,
        source_platform=source_platform,
        candidates=candidates,
        current_user=current_user,
        db=db,
    )
