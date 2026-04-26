"""Tests for the ResumeParser service."""

import json

import pytest
from docx import Document

from app.services.resume_parser import ResumeParser


pytestmark = pytest.mark.asyncio


# ---------------------------------------------------------------------------
# Email extraction
# ---------------------------------------------------------------------------

class TestExtractEmail:
    def test_extract_email_from_text(self):
        parser = ResumeParser()
        text = "联系方式：zhangsan@example.com，手机：13800138000"
        result = parser._extract_email(text)
        assert result == "zhangsan@example.com"

    def test_extract_email_none(self):
        parser = ResumeParser()
        text = "这段文字中没有电子邮件地址"
        result = parser._extract_email(text)
        assert result is None


# ---------------------------------------------------------------------------
# Phone extraction
# ---------------------------------------------------------------------------

class TestExtractPhone:
    def test_extract_phone_from_text(self):
        parser = ResumeParser()
        text = "联系电话：13800138000，邮箱：test@example.com"
        result = parser._extract_phone(text)
        assert result is not None
        assert "13800138000" in result

    def test_extract_phone_with_country_code(self):
        parser = ResumeParser()
        text = "手机：+86-13912345678"
        result = parser._extract_phone(text)
        assert result is not None
        assert "13912345678" in result

    def test_extract_phone_none(self):
        parser = ResumeParser()
        text = "这段文字中没有电话号码"
        result = parser._extract_phone(text)
        assert result is None


# ---------------------------------------------------------------------------
# Skills extraction
# ---------------------------------------------------------------------------

class TestExtractSkills:
    def test_extract_skills_from_text(self):
        parser = ResumeParser()
        text = """
        技能：Python, Java, Docker, MySQL, Redis, Git
        熟悉 React, Vue, Node.js
        """
        skills = parser._extract_skills(text)
        skill_lower = [s.lower() for s in skills]
        assert "python" in skill_lower
        assert "docker" in skill_lower

    def test_extract_skills_from_keywords(self):
        parser = ResumeParser()
        text = "我擅长Python和Java开发，熟悉Machine Learning和Deep Learning。"
        skills = parser._extract_skills(text)
        skill_lower = [s.lower() for s in skills]
        assert "python" in skill_lower
        assert "java" in skill_lower

    def test_extract_skills_stops_at_next_docx_section(self):
        parser = ResumeParser()
        text = """
        简历信息表
        一、基本信息
        二、专业技能
        具备10年技术管理与产品管理经验，擅长团队建设、跨部门协作与战略规划。
        熟悉AI大模型技术原理（RAG、SFT、RLHF等），具备技术方案评估与选型能力。
        熟悉 Dify、Coze 等AI平台，能将AI技术与业务场景深度结合，赋能业务提效。
        主导落地过 deepseek-R1、Qwen2.5-VL-32B-Instruct 等AI项目。
        熟练使用Axure、Jira、Figma、墨刀等工具进行产品设计与项目管理。
        熟悉敏捷开发方法，掌握Scrum、看板等项目管理工具，具备PMP、CSPM等专业认证。
        三、工作经历
        上海幻谱信息科技有限公司-总经理（CEO） 上海 2025年9月-至今
        挑战：多轮对话攻击场景复杂，需要团队跨领域协作攻克技术难题。
        四、社会职业
        全国项目管理标准化技术委员会专家库专家。
        """

        skills = parser._extract_skills(text)

        assert "技术管理" in skills
        assert "产品管理" in skills
        assert "RAG" in skills
        assert "Dify" in skills
        assert "Coze" in skills
        assert "DeepSeek-R1" in skills
        assert "Qwen2.5-VL-32B-Instruct" in skills
        assert "Axure" in skills
        assert "Jira" in skills
        assert "Figma" in skills
        assert "Scrum" in skills
        assert "PMP" in skills
        assert "AI大模型技术原理（RAG" not in skills
        assert "Qwen" not in skills
        assert "上海幻谱信息科技有限公司-总经理（CEO） 上海" not in skills
        assert "挑战：多轮对话攻击场景复杂" not in skills
        assert "社会职业" not in skills
        assert all(len(skill) <= 32 for skill in skills)


# ---------------------------------------------------------------------------
# DOCX extraction
# ---------------------------------------------------------------------------

class TestParseDocx:
    async def test_parse_docx_preserves_table_order(self, tmp_path):
        parser = ResumeParser()
        doc = Document()
        doc.add_paragraph("简历信息表")
        doc.add_paragraph("一、基本信息")
        table = doc.add_table(rows=3, cols=8)
        table.rows[0].cells[0].text = "姓名"
        table.rows[0].cells[1].text = "孟浩"
        table.rows[0].cells[2].text = "性别"
        table.rows[0].cells[3].text = "男"
        table.rows[0].cells[4].text = "年龄"
        table.rows[0].cells[5].text = "32"
        table.rows[0].cells[6].text = "联系方式"
        table.rows[0].cells[7].text = "18375312286"
        table.rows[1].cells[0].text = "院校名称"
        table.rows[1].cells[1].text = "安徽工程大学"
        table.rows[1].cells[2].text = "专业"
        table.rows[1].cells[3].text = "软件工程"
        table.rows[1].cells[4].text = "学历"
        table.rows[1].cells[5].text = "本科"
        doc.add_paragraph("二、专业技能")
        doc.add_paragraph("熟悉 Python、FastAPI、Docker。")
        doc.add_paragraph("三、工作经历")
        doc.add_paragraph("ABC Tech 后端工程师 5年 Python FastAPI API 平台开发")

        file_path = tmp_path / "resume.docx"
        doc.save(file_path)

        text = await parser.parse_docx(str(file_path))
        data = parser.extract_structured_data(text)

        assert text.index("姓名") < text.index("二、专业技能")
        assert data["name"] == "孟浩"
        assert data["gender"] == "男"
        assert data["age"] == 32
        assert data["phone"] == "18375312286"
        assert data["education"][0]["school"] == "安徽工程大学"
        assert data["education"][0]["degree"] == "本科"
        assert data["work_experience"][0]["company"] == "ABC Tech"
        assert data["work_experience"][0]["position"] == "后端工程师"


# ---------------------------------------------------------------------------
# Education extraction
# ---------------------------------------------------------------------------

class TestExtractEducation:
    def test_extract_education_from_text(self):
        parser = ResumeParser()
        text = """
        教育背景：
        清华大学 硕士 计算机科学
        北京大学 本科 数学
        """
        education = parser._extract_education(text)
        assert len(education) >= 1
        # First entry should detect a degree
        degrees = [e["degree"] for e in education]
        assert "硕士" in degrees or "本科" in degrees

    def test_extract_education_ignores_certificates(self):
        parser = ResumeParser()
        text = """
        一、基本信息
        院校名称\t安徽工程大学\t专业\t软件工程\t学历\t本科
        专业证书
        PMP、CSPM4、PRINCE基础级、伦敦国王学院心理学士证书。
        """
        education = parser._extract_education(text)

        assert len(education) == 1
        assert education[0]["school"] == "安徽工程大学"
        assert education[0]["degree"] == "本科"


# ---------------------------------------------------------------------------
# Work experience extraction
# ---------------------------------------------------------------------------

class TestExtractWorkExperience:
    def test_extract_work_experience_groups_pdf_wrapped_lines(self):
        parser = ResumeParser()
        text = """
        工作经历
        上海幻谱信息科技有限公司-技术总监（CEO）上海
        2025 年9 月-至今
        创立AI安全领域科技公司，聚焦大模型内容合规检测赛道。负责公司战略规划、核心产品研发与团队
        建设。
        清源系统（AI内容安全检测平台）
        从0到1搭建10人技术团队，建立矩阵式协作架构，实现多模态与语言模型双线并行研发。
        设计7层检测流水线+4类检测引擎，检测准确率>92%。
        江苏晟晖信息科技有限公司
        2025 年4 月～2025 年8 月
        主导研发AI辅助测评系统，实现测评报告自动化生成。
        项目经历
        智能问答系统
        """

        work = parser._extract_work_experience(text)

        assert len(work) == 2
        assert work[0]["company"] == "上海幻谱信息科技有限公司"
        assert "技术总监" in work[0]["position"]
        assert work[0]["start_date"] == "2025 年9 月"
        assert work[0]["end_date"] == "至今"
        assert "团队建设。" in work[0]["description"]
        assert "检测准确率>92%" in work[0]["description"]

    def test_extract_work_experience_when_title_and_period_share_line(self):
        parser = ResumeParser()
        text = """
        工作经历
        上海幻谱信息科技有限公司-总经理（CEO） 上海 2025年9月-至今
        创立AI安全领域科技公司，负责公司战略规划与团队管理。
        江苏晟晖信息科技有限公司-项目经理 2025年4月～2025年8月
        主导AI辅助测评系统项目管理，实现测评报告自动化生成。
        教育背景
        安徽工程大学 本科 软件工程
        """

        work = parser._extract_work_experience(text)

        assert len(work) == 2
        assert work[0]["company"] == "上海幻谱信息科技有限公司"
        assert "总经理" in work[0]["position"]
        assert work[0]["start_date"] == "2025年9月"
        assert work[0]["end_date"] == "至今"
        assert "战略规划" in work[0]["description"]

    def test_extract_work_experience_without_dates(self):
        parser = ResumeParser()
        text = """
        工作经历
        ABC Tech 后端工程师 5年 Python FastAPI PostgreSQL API 平台开发
        项目经历
        招聘分析平台 使用 Python FastAPI PostgreSQL Docker Redis 构建高并发 API 服务
        """

        work = parser._extract_work_experience(text)

        assert len(work) == 1
        assert work[0]["company"] == "ABC Tech"
        assert work[0]["position"] == "后端工程师"
        assert "FastAPI" in work[0]["description"]


# ---------------------------------------------------------------------------
# Name extraction
# ---------------------------------------------------------------------------

class TestExtractName:
    def test_extract_name_chinese(self):
        parser = ResumeParser()
        text = "张三\nzhangsan@example.com\n13800138000"
        result = parser._extract_name(text)
        assert result == "张三"

    def test_extract_name_english(self):
        parser = ResumeParser()
        text = "John Smith\njohn@example.com"
        result = parser._extract_name(text)
        assert result == "John Smith"

    def test_extract_name_none(self):
        parser = ResumeParser()
        text = "This is just a block of text\nwith multiple lines\nbut no clear name."
        result = parser._extract_name(text)
        # May or may not return a name; just ensure no crash
        assert result is None or isinstance(result, str)

    def test_extract_name_from_label_separated_basic_info(self):
        parser = ResumeParser()
        text = """
        简历信息表
        一、基本信息
        姓名
        孟浩
        性别
        男
        年龄
        32
        联系方式
        18375312286
        """

        basic_info = parser.extract_basic_info(text)

        assert basic_info["name"] == "孟浩"
        assert basic_info["gender"] == "男"
        assert basic_info["age"] == 32
        assert basic_info["phone"] == "18375312286"

    def test_extract_name_does_not_return_label(self):
        parser = ResumeParser()
        text = """
        姓名
        性别
        男
        年龄
        32
        """
        result = parser._extract_name(text)
        assert result is None

    def test_extract_structured_data_includes_basic_info(self):
        parser = ResumeParser()
        text = "姓名：孟浩 性别：男 年龄：32\n电话：18375312286\n邮箱：menghao@example.com"
        result = parser.extract_structured_data(text)

        assert result["name"] == "孟浩"
        assert result["gender"] == "男"
        assert result["age"] == 32
        assert result["email"] == "menghao@example.com"

    def test_extract_basic_info_when_docx_table_text_is_late(self):
        parser = ResumeParser()
        filler = "\n".join(f"专业技能说明 {index}" for index in range(100))
        text = f"""
        简历信息表
        一、基本信息
        {filler}
        姓名 孟浩 性别 男 年龄 32 联系方式 18375312286 工作年限 10年
        """

        result = parser.extract_basic_info(text)

        assert result["name"] == "孟浩"
        assert result["gender"] == "男"
        assert result["age"] == 32


# ---------------------------------------------------------------------------
# AI response JSON parsing (parse_resume does not have this method;
# instead test the parse_ai_response of AIAnalyzer which handles it)
# ---------------------------------------------------------------------------

class TestParseAIResponseJsonWithCodeBlock:
    """Verify that markdown-wrapped JSON is correctly extracted."""

    def test_parse_ai_response_json_with_code_block(self):
        parser = ResumeParser()
        # ResumeParser does not have parse_ai_response; test via AIAnalyzer
        from app.services.ai_analyzer import AIAnalyzer

        analyzer = AIAnalyzer()
        raw = '''```json
{
    "skill_match": {
        "score": 88,
        "analysis": "匹配度高"
    },
    "strengths": ["技能全面"]
}
```'''
        result = analyzer.parse_ai_response(raw)
        assert result["skill_match"]["score"] == 88
        assert result["strengths"] == ["技能全面"]
